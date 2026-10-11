"""
Build DOCX and PDF versions of report/report.md.

Usage:
    python scripts/build_report_docs.py

Outputs:
    report/report.docx
    report/report.pdf
"""

import os
import re
import sys
import subprocess
from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_MD = BASE_DIR / "report" / "report.md"
OUTPUT_DOCX = BASE_DIR / "report" / "report.docx"
OUTPUT_PDF = BASE_DIR / "report" / "report.pdf"
CONVERT_PS1 = BASE_DIR / "scripts" / "docx_to_pdf.ps1"

COLOR_PRIMARY = RGBColor(0x1A, 0x36, 0x5D)    # Deep Navy
COLOR_SECONDARY = RGBColor(0x2B, 0x6C, 0xB0)  # Medium Blue
COLOR_TEXT = RGBColor(0x2D, 0x37, 0x48)       # Charcoal Body Text
COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)      # Muted Gray
COLOR_HEADER_BG = "EDF2F7"                    # Light Slate Table Header

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_formatted_runs(paragraph, text, base_font_size=10.5, is_italic=False, is_bold=False, color=COLOR_TEXT):
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = "Calibri"
        run.font.size = Pt(base_font_size)
        run.font.color.rgb = color
        
        if token.startswith('***') and token.endswith('***') and len(token) >= 6:
            run.text = token[3:-3]
            run.bold = True
            run.italic = True
        elif token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.text = token[2:-2]
            run.bold = True
            run.italic = is_italic
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.text = token[1:-1]
            run.bold = is_bold
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_font_size * 0.95)
            run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
        else:
            run.text = token
            run.bold = is_bold
            run.italic = is_italic

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Medical Image Compression | Final Project Report")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(9)
        f_run.font.color.rgb = COLOR_MUTED
    
    yaml_meta = {}
    lines = content.splitlines()
    i = 0
    if lines and lines[0].strip() == '---':
        i = 1
        while i < len(lines) and lines[i].strip() != '---':
            line = lines[i]
            if ':' in line:
                k, v = line.split(':', 1)
                yaml_meta[k.strip()] = v.strip().strip('"').strip("'")
            i += 1
        i += 1
    
    title = yaml_meta.get('title', 'Medical Image Compression Project Report')
    author = yaml_meta.get('author', 'Medical Image Compression Team')
    date = yaml_meta.get('date', '2026')
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(8)
    run_t = p_title.add_run(title)
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(22)
    run_t.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(24)
    run_m = p_meta.add_run(f"{author}  |  {date}  |  Chest Radiograph Compression Benchmark")
    run_m.font.name = "Calibri"
    run_m.font.size = Pt(10.5)
    run_m.font.color.rgb = COLOR_MUTED
    
    n = len(lines)
    while i < n:
        raw_line = lines[i]
        line = raw_line.strip()
        
        if not line:
            i += 1
            continue
        
        if line in ('---', '***', '___'):
            i += 1
            continue
        
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            title_text = line.lstrip('#').strip()
            
            p_head = doc.add_paragraph()
            p_head.paragraph_format.keep_with_next = True
            
            if level == 1:
                p_head.paragraph_format.space_before = Pt(18)
                p_head.paragraph_format.space_after = Pt(6)
                run = p_head.add_run(title_text)
                run.font.name = "Calibri"
                run.font.size = Pt(16)
                run.bold = True
                run.font.color.rgb = COLOR_PRIMARY
            elif level == 2:
                p_head.paragraph_format.space_before = Pt(14)
                p_head.paragraph_format.space_after = Pt(4)
                run = p_head.add_run(title_text)
                run.font.name = "Calibri"
                run.font.size = Pt(13)
                run.bold = True
                run.font.color.rgb = COLOR_SECONDARY
            elif level == 3:
                p_head.paragraph_format.space_before = Pt(10)
                p_head.paragraph_format.space_after = Pt(3)
                run = p_head.add_run(title_text)
                run.font.name = "Calibri"
                run.font.size = Pt(11.5)
                run.bold = True
                run.font.color.rgb = COLOR_TEXT
            else:
                p_head.paragraph_format.space_before = Pt(8)
                p_head.paragraph_format.space_after = Pt(2)
                run = p_head.add_run(title_text)
                run.font.name = "Calibri"
                run.font.size = Pt(10.5)
                run.bold = True
                run.italic = True
                run.font.color.rgb = COLOR_MUTED
            
            i += 1
            continue
        
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            caption, img_rel_path = img_match.groups()
            img_full_path = (REPORT_MD.parent / img_rel_path).resolve()
            if not img_full_path.exists():
                img_full_path = (BASE_DIR / img_rel_path).resolve()
            
            if img_full_path.exists():
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.paragraph_format.keep_with_next = True
                run_pic = p_img.add_run()
                run_pic.add_picture(str(img_full_path), width=Inches(5.6))
            else:
                print(f"Warning: image path not found: {img_full_path}")
            
            i += 1
            if i < n and (lines[i].strip().startswith('**Figure') or lines[i].strip().startswith('*Figure')):
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                add_formatted_runs(p_cap, lines[i].strip(), base_font_size=9.5, is_italic=True, color=COLOR_MUTED)
                i += 1
            continue
        
        if line.startswith('**Table') and i + 1 < n and '|' in lines[i + 1]:
            p_tbl_title = doc.add_paragraph()
            p_tbl_title.paragraph_format.space_before = Pt(12)
            p_tbl_title.paragraph_format.space_after = Pt(3)
            p_tbl_title.paragraph_format.keep_with_next = True
            add_formatted_runs(p_tbl_title, line, base_font_size=10.5, is_bold=True, color=COLOR_PRIMARY)
            i += 1
            continue
        
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                parsed_rows = []
                alignments = []
                for idx, t_line in enumerate(table_lines):
                    cells = [c.strip() for c in t_line.strip('|').split('|')]
                    if idx == 1 and all(set(c).issubset({'-', ':', ' '}) for c in cells):
                        for c in cells:
                            if c.startswith(':') and c.endswith(':'):
                                alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
                            elif c.endswith(':'):
                                alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
                            else:
                                alignments.append(WD_ALIGN_PARAGRAPH.LEFT)
                        continue
                    parsed_rows.append(cells)
                
                if parsed_rows:
                    num_cols = max(len(r) for r in parsed_rows)
                    num_rows = len(parsed_rows)
                    
                    tbl = doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.style = 'Table Grid'
                    
                    for r_idx, row in enumerate(parsed_rows):
                        is_header = (r_idx == 0)
                        tr = tbl.rows[r_idx]._tr
                        trPr = tr.get_or_add_trPr()
                        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                        if is_header:
                            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                            
                        for c_idx in range(num_cols):
                            cell_text = row[c_idx] if c_idx < len(row) else ""
                            cell = tbl.cell(r_idx, c_idx)
                            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                            
                            if is_header:
                                set_cell_background(cell, COLOR_HEADER_BG)
                            
                            p_cell = cell.paragraphs[0]
                            p_cell.paragraph_format.space_before = Pt(2)
                            p_cell.paragraph_format.space_after = Pt(2)
                            
                            align = alignments[c_idx] if c_idx < len(alignments) else WD_ALIGN_PARAGRAPH.LEFT
                            p_cell.alignment = align
                            
                            add_formatted_runs(
                                p_cell, 
                                cell_text, 
                                base_font_size=9.0, 
                                is_bold=is_header, 
                                color=COLOR_PRIMARY if is_header else COLOR_TEXT
                            )
                    
                    p_spacer = doc.add_paragraph()
                    p_spacer.paragraph_format.space_before = Pt(0)
                    p_spacer.paragraph_format.space_after = Pt(6)
            continue
        
        bullet_match = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', raw_line)
        if bullet_match:
            indent_spaces, marker, item_text = bullet_match.groups()
            p_list = doc.add_paragraph()
            p_list.paragraph_format.space_before = Pt(1)
            p_list.paragraph_format.space_after = Pt(3)
            
            indent_level = len(indent_spaces) // 2
            p_list.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            
            run_sym = p_list.add_run(f"{marker} " if marker.endswith('.') else "• ")
            run_sym.font.name = "Calibri"
            run_sym.font.size = Pt(10.5)
            run_sym.bold = True
            run_sym.font.color.rgb = COLOR_SECONDARY
            
            add_formatted_runs(p_list, item_text, base_font_size=10.5, color=COLOR_TEXT)
            i += 1
            continue
        
        p_body = doc.add_paragraph()
        p_body.paragraph_format.space_before = Pt(2)
        p_body.paragraph_format.space_after = Pt(6)
        p_body.paragraph_format.line_spacing = 1.15
        
        if line.startswith('>'):
            p_body.paragraph_format.left_indent = Inches(0.4)
            add_formatted_runs(p_body, line.lstrip('>').strip(), base_font_size=10.0, is_italic=True, color=COLOR_MUTED)
        else:
            add_formatted_runs(p_body, line, base_font_size=10.5, color=COLOR_TEXT)
        
        i += 1

    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path} ({os.path.getsize(docx_path):,} bytes)")

def convert_docx_to_pdf(docx_path, pdf_path):
    if not CONVERT_PS1.exists():
        print(f"Error: {CONVERT_PS1} not found")
        return
    res = subprocess.run(
        ["powershell", "-ExecutionPolicy", "Bypass", "-File", str(CONVERT_PS1)],
        capture_output=True,
        text=True
    )
    print(res.stdout)
    if res.stderr:
        print("PowerShell Stderr:", res.stderr)

if __name__ == "__main__":
    print(f"Building report documents from {REPORT_MD} ...")
    parse_markdown_to_docx(REPORT_MD, OUTPUT_DOCX)
    print("Converting to PDF via MS Word COM...")
    convert_docx_to_pdf(OUTPUT_DOCX, OUTPUT_PDF)
